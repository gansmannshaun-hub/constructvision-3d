"""Materials takeoff PDF/CSV/XLSX generation."""
from __future__ import annotations

import csv
import re
from collections import defaultdict
from datetime import datetime, timezone
from io import BytesIO, StringIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import billing as billing_mod
from pricing import compute_bid, get_pricing_config
from routes.projects import get_or_create_blueprint
from utilities_takeoff import compute_utilities_takeoff


async def _build_combined_materials(db, project_id: str) -> list[dict]:
    """Materials from DB + auto-computed utilities, in display order."""
    mats = await db.materials.find({"project_id": project_id}, {"_id": 0}).sort("category", 1).to_list(2000)
    bp = await get_or_create_blueprint(db, project_id)
    utilities = compute_utilities_takeoff(bp)
    return list(mats) + utilities


def _safe_name(s: str) -> str:
    return re.sub(r"[^a-zA-Z0-9_-]+", "_", s or "")[:40] or "project"


def _src_label(m: dict) -> str:
    if m.get("auto_computed"):
        return "AUTO"
    if m.get("ai_extracted"):
        return "AI"
    return ""


def build_takeoff_router(db, get_current_user) -> APIRouter:
    router = APIRouter(prefix="/api")

    @router.get("/projects/{project_id}/takeoff.pdf")
    async def takeoff_pdf(project_id: str, user: dict = Depends(get_current_user)):
        proj_check = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 1})
        if not proj_check:
            raise HTTPException(404, "Project not found")

        user = await billing_mod.ensure_user_subscription(db, user)
        ok, reason = await billing_mod.can_download_pdf(db, user)
        if not ok:
            raise HTTPException(402, reason)

        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
        )
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch

        proj = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
        if not proj:
            raise HTTPException(404, "Project not found")
        await billing_mod.incr_usage(db, user["id"], "pdf_downloads")
        await db.users.update_one(
            {"id": user["id"]}, {"$inc": {"entitlements.lifetime_pdfs": 1}}
        )
        premium_branding = bool((user.get("entitlements") or {}).get("pdf_premium_branding"))

        mats = await db.materials.find({"project_id": project_id}, {"_id": 0}).sort("category", 1).to_list(1000)
        docs = await db.documents.find(
            {"project_id": project_id, "status": "done"}, {"_id": 0, "image_base64": 0}
        ).sort("created_at", -1).to_list(50)

        # Compute auto-utilities from the live blueprint and merge them in
        bp = await get_or_create_blueprint(db, project_id)
        utilities = compute_utilities_takeoff(bp)
        mats_combined = list(mats) + utilities

        buf = BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=LETTER,
            leftMargin=0.6 * inch, rightMargin=0.6 * inch,
            topMargin=0.6 * inch, bottomMargin=0.6 * inch,
            title=f"Takeoff — {proj['name']}",
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("Title", parent=styles["Title"], fontName="Helvetica-Bold",
            fontSize=28, textColor=colors.HexColor("#0A0A0A"), spaceAfter=4, leading=30)
        sub_style = ParagraphStyle("Sub", parent=styles["Normal"], fontName="Helvetica",
            fontSize=10, textColor=colors.HexColor("#666666"), spaceAfter=18)
        label_style = ParagraphStyle("Lbl", parent=styles["Normal"], fontName="Helvetica-Bold",
            fontSize=8, textColor=colors.HexColor("#888888"), spaceAfter=2)
        h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontName="Helvetica-Bold",
            fontSize=14, textColor=colors.HexColor("#0A0A0A"), spaceBefore=10, spaceAfter=8)

        elements = []
        elements.append(Paragraph("ATLAS&nbsp;&nbsp;<font color='#888888'>// CONSTRUCTION OS</font>", label_style))
        elements.append(Paragraph("Material Takeoff", title_style))
        elements.append(Paragraph(
            f"Project: <b>{proj['name']}</b> &nbsp;·&nbsp; "
            f"Generated: {datetime.now(timezone.utc).strftime('%b %d, %Y %H:%M UTC')}"
            + (" &nbsp;·&nbsp; <font color='#FFCC00'><b>PREMIUM</b></font>" if premium_branding else ""),
            sub_style,
        ))
        if premium_branding:
            elements.append(Paragraph(
                f"<font color='#0055FF' size='11'><b>Prepared by {user.get('name', '')} · {user.get('email', '')}</b></font>",
                sub_style,
            ))

        grouped = defaultdict(list)
        for m in mats_combined:
            grouped[m.get("category") or "Other"].append(m)

        grand_total = 0.0
        cat_subtotals = {}
        for cat, items in grouped.items():
            s = sum(float(i.get("quantity") or 0) * float(i.get("unit_price") or 0) for i in items)
            cat_subtotals[cat] = s
            grand_total += s

        auto_count = sum(1 for m in mats_combined if m.get("auto_computed"))
        ai_count = sum(1 for m in mats_combined if m.get("ai_extracted"))

        # KPI bar
        kpi_data = [[
            Paragraph("<b>TOTAL ITEMS</b>", label_style),
            Paragraph("<b>CATEGORIES</b>", label_style),
            Paragraph("<b>AI / AUTO</b>", label_style),
            Paragraph("<b>PROJECT COST</b>", label_style),
        ], [
            Paragraph(f"<font size='18'><b>{len(mats_combined)}</b></font>", styles["Normal"]),
            Paragraph(f"<font size='18'><b>{len(grouped)}</b></font>", styles["Normal"]),
            Paragraph(
                f"<font size='14'><b>{ai_count}</b></font>"
                f" <font color='#888888' size='10'>AI</font>"
                f" &nbsp;<font size='14' color='#FF6600'><b>{auto_count}</b></font>"
                f" <font color='#888888' size='10'>AUTO</font>",
                styles["Normal"],
            ),
            Paragraph(f"<font size='18' color='#0055FF'><b>${grand_total:,.2f}</b></font>", styles["Normal"]),
        ]]
        kpi = Table(kpi_data, colWidths=[1.7 * inch] * 4)
        kpi.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#EEEEEE")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F7F7F7")),
            ("PADDING", (0, 0), (-1, -1), 10),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(kpi)
        elements.append(Spacer(1, 14))

        if utilities:
            elements.append(Paragraph(
                "<font color='#FF6600' size='9'><b>// AUTO-COMPUTED FROM BLUEPRINT</b></font>"
                f" &nbsp;<font color='#666666' size='9'>{len(utilities)} utilities &amp; MEP line items "
                f"derived procedurally from the wall footprint — septic, underground, plumbing rough-in, electrical.</font>",
                sub_style,
            ))

        if not mats_combined:
            elements.append(Paragraph(
                "<i>No materials in this project yet. Upload a blueprint to auto-populate.</i>",
                sub_style,
            ))
        else:
            for cat in sorted(grouped.keys()):
                items = grouped[cat]
                elements.append(Paragraph(
                    f"{cat} &nbsp;<font color='#888888' size='10'>· {len(items)} items · "
                    f"${cat_subtotals[cat]:,.2f}</font>", h2,
                ))
                table_data = [["#", "Material", "Qty", "Unit", "Unit Price", "Line Total", "Src"]]
                for idx, m in enumerate(items, 1):
                    qty = float(m.get("quantity") or 0)
                    price = float(m.get("unit_price") or 0)
                    line_total = qty * price
                    if m.get("auto_computed"):
                        src = "AUTO"
                    elif m.get("ai_extracted"):
                        src = "AI"
                    else:
                        src = ""
                    table_data.append([
                        str(idx),
                        m.get("name", "")[:48],
                        f"{qty:g}",
                        m.get("unit", ""),
                        f"${price:,.2f}",
                        f"${line_total:,.2f}",
                        src,
                    ])
                table_data.append(["", "", "", "", "Subtotal", f"${cat_subtotals[cat]:,.2f}", ""])
                t = Table(table_data, colWidths=[
                    0.3 * inch, 2.7 * inch, 0.65 * inch, 0.65 * inch,
                    0.9 * inch, 1.0 * inch, 0.45 * inch,
                ])
                style = [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A0A0A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#FFCC00")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 8),
                    ("FONTSIZE", (0, 1), (-1, -1), 9),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F7F7F7")),
                    ("ALIGN", (2, 1), (5, -1), "RIGHT"),
                    ("ALIGN", (6, 1), (6, -1), "CENTER"),
                    ("ALIGN", (0, 0), (-1, 0), "LEFT"),
                    ("ALIGN", (4, 0), (5, 0), "RIGHT"),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDDDDD")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#FAFAFA")]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("FONTSIZE", (6, 1), (6, -2), 7),
                ]
                # Color-code Src column
                for row_idx, m in enumerate(items, start=1):
                    if m.get("auto_computed"):
                        style.append(("TEXTCOLOR", (6, row_idx), (6, row_idx), colors.HexColor("#FF6600")))
                        style.append(("FONTNAME", (6, row_idx), (6, row_idx), "Helvetica-Bold"))
                    elif m.get("ai_extracted"):
                        style.append(("TEXTCOLOR", (6, row_idx), (6, row_idx), colors.HexColor("#0055FF")))
                        style.append(("FONTNAME", (6, row_idx), (6, row_idx), "Helvetica-Bold"))
                t.setStyle(TableStyle(style))
                elements.append(t)
                elements.append(Spacer(1, 14))

            gt = Table([["GRAND TOTAL", f"${grand_total:,.2f}"]], colWidths=[5.0 * inch, 1.55 * inch])
            gt.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#0055FF")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 14),
                ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 14),
                ("RIGHTPADDING", (0, 0), (-1, -1), 14),
                ("TOPPADDING", (0, 0), (-1, -1), 12),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
            ]))
            elements.append(gt)

        if docs:
            elements.append(PageBreak())
            elements.append(Paragraph("Source Documents", h2))
            for d in docs:
                summary = (d.get("analysis") or {}).get("summary") or ""
                row = (
                    f"<b>{d.get('filename', '')}</b> "
                    f"<font color='#888888' size='8'>· {d.get('doc_type') or 'unknown'} · "
                    f"{d.get('materials_count') or 0} materials"
                    f"{' · 3D synced' if d.get('synced_3d') else ''}</font>"
                )
                elements.append(Paragraph(row, styles["Normal"]))
                if summary:
                    elements.append(Paragraph(
                        f"<font color='#555555' size='9'>{summary}</font>",
                        styles["Normal"],
                    ))
                elements.append(Spacer(1, 10))

        elements.append(Spacer(1, 20))
        elements.append(Paragraph(
            "<font color='#888888' size='8'><i>Prices are AI-estimated US 2026 trade rates "
            "(AI = vision-extracted from uploaded docs; AUTO = procedurally derived from blueprint footprint). "
            "Verify with vendors before final bidding. Generated by Atlas Construction OS.</i></font>",
            styles["Normal"],
        ))

        doc.build(elements)
        buf.seek(0)
        safe_name = re.sub(r"[^a-zA-Z0-9_-]+", "_", proj["name"])[:40] or "project"
        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="atlas_takeoff_{safe_name}.pdf"',
            },
        )

    @router.get("/projects/{project_id}/takeoff.csv")
    async def takeoff_csv(project_id: str, user: dict = Depends(get_current_user)):
        proj = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
        if not proj:
            raise HTTPException(404, "Project not found")
        mats = await _build_combined_materials(db, project_id)
        cfg = await get_pricing_config(db, project_id)
        result = compute_bid(mats, cfg)
        totals = result["totals"]
        sio = StringIO()
        writer = csv.writer(sio)
        writer.writerow([
            "Category", "Material", "Quantity", "Unit",
            "Material $/unit", "Labor $/unit", "Material Total", "Labor Total",
            "Line Total", "Source",
        ])
        for m in sorted(result["lines"], key=lambda x: (x.get("category") or "", x.get("name") or "")):
            writer.writerow([
                m.get("category") or "Other",
                m.get("name") or "",
                f"{float(m.get('quantity') or 0):g}",
                m.get("unit") or "",
                f"{float(m.get('unit_price') or 0):.2f}",
                f"{float(m.get('labor_unit_price') or 0):.2f}",
                f"{m.get('material_total', 0):.2f}",
                f"{m.get('labor_total', 0):.2f}",
                f"{m.get('line_total', 0):.2f}",
                _src_label(m),
            ])
        writer.writerow([])
        loc = f"{cfg.get('city') or cfg.get('state') or ''} ({cfg.get('zip') or 'no ZIP'})".strip()
        writer.writerow(["", f"Regional multiplier ({loc})", "", "", "", "", "", "", f"×{totals['applied_multiplier']:.3f}", ""])
        writer.writerow(["", "Materials subtotal", "", "", "", "", "", "", f"{totals['materials_subtotal']:.2f}", ""])
        writer.writerow(["", "Labor subtotal",     "", "", "", "", "", "", f"{totals['labor_subtotal']:.2f}", ""])
        writer.writerow(["", "Base subtotal",      "", "", "", "", "", "", f"{totals['base_subtotal']:.2f}", ""])
        writer.writerow(["", f"Waste ({totals['waste_pct']}%)",       "", "", "", "", "", "", f"{totals['waste_amount']:.2f}", ""])
        writer.writerow(["", f"Overhead ({totals['overhead_pct']}%)", "", "", "", "", "", "", f"{totals['overhead_amount']:.2f}", ""])
        writer.writerow(["", f"Profit ({totals['profit_pct']}%)",     "", "", "", "", "", "", f"{totals['profit_amount']:.2f}", ""])
        writer.writerow(["", f"Contingency ({totals['contingency_pct']}%)", "", "", "", "", "", "", f"{totals['contingency_amount']:.2f}", ""])
        writer.writerow(["", "GRAND TOTAL",        "", "", "", "", "", "", f"{totals['grand_total']:.2f}", ""])
        out = sio.getvalue().encode("utf-8-sig")
        safe = _safe_name(proj["name"])
        return StreamingResponse(
            BytesIO(out),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="atlas_takeoff_{safe}.csv"'},
        )

    @router.get("/projects/{project_id}/takeoff.xlsx")
    async def takeoff_xlsx(project_id: str, user: dict = Depends(get_current_user)):
        proj = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
        if not proj:
            raise HTTPException(404, "Project not found")
        mats = await _build_combined_materials(db, project_id)
        cfg = await get_pricing_config(db, project_id)
        result = compute_bid(mats, cfg)
        totals = result["totals"]

        wb = Workbook()
        ws = wb.active
        ws.title = "Takeoff"

        ws["A1"] = f"Atlas Takeoff — {proj['name']}"
        ws["A1"].font = Font(name="Calibri", size=16, bold=True)
        ws["A2"] = f"Generated {datetime.now(timezone.utc).strftime('%b %d, %Y %H:%M UTC')}"
        ws["A2"].font = Font(name="Calibri", size=10, color="666666")
        loc = f"{cfg.get('city') or cfg.get('state') or ''} {cfg.get('zip') or ''}".strip()
        ws["A3"] = f"Region: {loc or 'US average'} · multiplier ×{totals['applied_multiplier']:.3f}"
        ws["A3"].font = Font(name="Calibri", size=10, color="0055FF", bold=True)
        ws.append([])

        header = ["#", "Category", "Material", "Quantity", "Unit",
                  "Material $/unit", "Labor $/unit", "Material Total", "Labor Total",
                  "Line Total", "Source"]
        ws.append(header)
        header_row = ws.max_row
        head_fill = PatternFill("solid", fgColor="0A0A0A")
        head_font = Font(name="Calibri", bold=True, color="FFCC00", size=11)
        for col in range(1, len(header) + 1):
            c = ws.cell(row=header_row, column=col)
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(horizontal="left", vertical="center")

        grouped: dict[str, list[dict]] = defaultdict(list)
        for m in result["lines"]:
            grouped[m.get("category") or "Other"].append(m)

        idx = 0
        for cat in sorted(grouped.keys()):
            for m in grouped[cat]:
                idx += 1
                qty = float(m.get("quantity") or 0)
                mat_p = float(m.get("unit_price") or 0)
                lab_p = float(m.get("labor_unit_price") or 0)
                ws.append([idx, cat, m.get("name") or "", qty, m.get("unit") or "",
                           mat_p, lab_p,
                           m.get("material_total", 0), m.get("labor_total", 0), m.get("line_total", 0),
                           _src_label(m)])
                row = ws.max_row
                for col in (6, 7, 8, 9, 10):
                    ws.cell(row=row, column=col).number_format = '"$"#,##0.00'
                src_cell = ws.cell(row=row, column=11)
                if src_cell.value == "AUTO":
                    src_cell.font = Font(bold=True, color="FF6600")
                elif src_cell.value == "AI":
                    src_cell.font = Font(bold=True, color="0055FF")

        ws.append([])
        bold = Font(bold=True, size=11)
        big = Font(bold=True, size=13, color="0055FF")

        def add_summary(label, amount, big_font=False):
            ws.append(["", "", "", "", "", "", "", "", label, amount, ""])
            r = ws.max_row
            ws.cell(row=r, column=9).font = big if big_font else bold
            ws.cell(row=r, column=10).number_format = '"$"#,##0.00'
            ws.cell(row=r, column=10).font = big if big_font else bold

        add_summary("Materials subtotal", totals["materials_subtotal"])
        add_summary("Labor subtotal",     totals["labor_subtotal"])
        add_summary("Base subtotal",      totals["base_subtotal"])
        add_summary(f"Waste ({totals['waste_pct']}%)",       totals["waste_amount"])
        add_summary(f"Overhead ({totals['overhead_pct']}%)", totals["overhead_amount"])
        add_summary(f"Profit ({totals['profit_pct']}%)",     totals["profit_amount"])
        add_summary(f"Contingency ({totals['contingency_pct']}%)", totals["contingency_amount"])
        add_summary("GRAND TOTAL",        totals["grand_total"], big_font=True)

        widths = [5, 16, 38, 10, 8, 12, 12, 13, 13, 13, 8]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A6"

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        safe = _safe_name(proj["name"])
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="atlas_takeoff_{safe}.xlsx"'},
        )

    return router
