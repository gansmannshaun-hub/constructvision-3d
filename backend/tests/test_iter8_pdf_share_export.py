"""Iteration 8 tests: multi-page PDF upload, CSV/XLSX export, share link feature."""
import io
import time
import uuid

import pytest
import requests
from openpyxl import load_workbook
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas


# ---------- Helpers ----------
def _make_2_page_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=LETTER)
    # Page 1 - Plan
    c.rect(80, 80, 450, 600, stroke=1, fill=0)
    c.rect(100, 400, 200, 250, stroke=1, fill=0)
    c.rect(310, 400, 210, 250, stroke=1, fill=0)
    c.setFont("Helvetica-Bold", 24)
    c.drawString(220, 720, "Plan")
    c.drawString(150, 520, "BEDROOM")
    c.drawString(360, 520, "KITCHEN")
    c.showPage()
    # Page 2 - Elevation
    c.rect(80, 200, 450, 350, stroke=1, fill=0)
    c.rect(120, 250, 80, 120, stroke=1, fill=0)
    c.setFont("Helvetica-Bold", 24)
    c.drawString(200, 720, "Elevation")
    c.showPage()
    c.save()
    return buf.getvalue()


def _make_project(base_url, auth_headers, name=None):
    name = name or f"TEST_iter8_{uuid.uuid4().hex[:6]}"
    r = requests.post(f"{base_url}/api/projects", json={"name": name, "description": "iter8"},
                      headers=auth_headers, timeout=20)
    assert r.status_code == 200, r.text
    return r.json()


def _wait_for_doc_done(base_url, auth_headers, project_id, timeout=180):
    deadline = time.time() + timeout
    last = None
    while time.time() < deadline:
        r = requests.get(f"{base_url}/api/projects/{project_id}/documents",
                         headers=auth_headers, timeout=20)
        if r.status_code == 200 and r.json():
            last = r.json()[0]
            if last.get("status") in ("done", "error"):
                return last
        time.sleep(3)
    return last


# ---------- Multi-page PDF upload ----------
class TestPdfUpload:
    def test_2_page_pdf_upload_and_pipeline(self, base_url, auth_headers_a):
        proj = _make_project(base_url, auth_headers_a)
        pdf_bytes = _make_2_page_pdf()
        files = {"file": ("test_plan.pdf", pdf_bytes, "application/pdf")}
        r = requests.post(f"{base_url}/api/projects/{proj['id']}/documents/upload",
                          files=files, headers=auth_headers_a, timeout=60)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("is_pdf") is True
        assert body.get("pages_total") == 2

        final = _wait_for_doc_done(base_url, auth_headers_a, proj["id"], timeout=240)
        assert final is not None, "no document record found"
        assert final["status"] == "done", f"status={final.get('status')} err={final.get('error')}"
        assert final.get("is_pdf") is True
        assert final.get("pages_total") == 2
        assert final.get("pages_done") == 2
        assert "materials_count" in final
        assert "materials_merged" in final
        assert "materials_skipped" in final

    def test_image_upload_still_works(self, base_url, auth_headers_a, floor_plan_png):
        proj = _make_project(base_url, auth_headers_a)
        with open(floor_plan_png, "rb") as fh:
            files = {"file": ("floor.png", fh.read(), "image/png")}
        r = requests.post(f"{base_url}/api/projects/{proj['id']}/documents/upload",
                          files=files, headers=auth_headers_a, timeout=60)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body.get("is_pdf") is False
        assert body.get("pages_total") == 1

    def test_oversize_file_rejected(self, base_url, auth_headers_a):
        proj = _make_project(base_url, auth_headers_a)
        # 17MB payload
        files = {"file": ("big.pdf", b"\0" * (17 * 1024 * 1024), "application/pdf")}
        r = requests.post(f"{base_url}/api/projects/{proj['id']}/documents/upload",
                          files=files, headers=auth_headers_a, timeout=60)
        assert r.status_code == 400
        assert "16MB" in r.text or "too large" in r.text.lower()

    def test_unsupported_filetype_rejected(self, base_url, auth_headers_a):
        proj = _make_project(base_url, auth_headers_a)
        files = {"file": ("notes.txt", b"hello", "text/plain")}
        r = requests.post(f"{base_url}/api/projects/{proj['id']}/documents/upload",
                          files=files, headers=auth_headers_a, timeout=30)
        assert r.status_code == 400
        assert "unsupported" in r.text.lower() or "pdf" in r.text.lower()


# ---------- CSV / XLSX export ----------
class TestExports:
    @pytest.fixture(scope="class")
    def project_with_blueprint(self, base_url, user_a):
        # Create a project; blueprint auto-utility computation happens regardless of materials
        headers = {"Authorization": f"Bearer {user_a['token']}"}
        proj = _make_project(base_url, headers)
        # Seed a few walls so utilities have something to compute on
        bp_payload = {
            "walls": [
                {"id": "w1", "start": [0, 0], "end": [40, 0], "thickness": 0.2},
                {"id": "w2", "start": [40, 0], "end": [40, 30], "thickness": 0.2},
                {"id": "w3", "start": [40, 30], "end": [0, 30], "thickness": 0.2},
                {"id": "w4", "start": [0, 30], "end": [0, 0], "thickness": 0.2},
            ],
            "doors": [], "windows": [], "labels": [],
            "roof_type": "gable", "roof_pitch_deg": 30,
        }
        r = requests.put(f"{base_url}/api/projects/{proj['id']}/blueprint",
                         json=bp_payload, headers=headers, timeout=20)
        assert r.status_code == 200, r.text
        return proj, headers

    def test_csv_export_with_bom_and_grand_total(self, base_url, project_with_blueprint):
        proj, headers = project_with_blueprint
        r = requests.get(f"{base_url}/api/projects/{proj['id']}/takeoff.csv",
                         headers=headers, timeout=30)
        assert r.status_code == 200, r.text
        assert "text/csv" in r.headers.get("content-type", "")
        body = r.content
        # UTF-8 BOM
        assert body[:3] == b"\xef\xbb\xbf"
        text = body.decode("utf-8-sig")
        first_line = text.splitlines()[0]
        assert first_line.startswith("Category,Material,Quantity,Unit,Material $/unit,Labor $/unit,Material Total,Labor Total,Line Total,Source")
        # Source labels present (AUTO since utilities are auto-computed)
        assert "AUTO" in text
        # Utilities present
        assert ("Septic" in text or "EMT" in text or "Water main" in text.lower() or "water main" in text.lower())
        # Last non-empty row should be GRAND TOTAL
        non_empty = [line for line in text.splitlines() if line.strip()]
        assert "GRAND TOTAL" in non_empty[-1]

    def test_xlsx_export_format(self, base_url, project_with_blueprint):
        proj, headers = project_with_blueprint
        r = requests.get(f"{base_url}/api/projects/{proj['id']}/takeoff.xlsx",
                         headers=headers, timeout=30)
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # Find header row (contains 'Material')
        header_row_idx = None
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=False):
            vals = [c.value for c in row]
            if "Material" in vals:
                header_row_idx = row[0].row
                # Material expected at column C (index 3, 1-based)
                assert vals[2] == "Material", f"Material at col {vals.index('Material')+1}, expected col 3"
                break
        assert header_row_idx is not None, "Header row not found"

        # Last non-empty row contains GRAND TOTAL
        last_vals = None
        for row in ws.iter_rows(values_only=True):
            if any(v not in (None, "") for v in row):
                last_vals = row
        assert last_vals is not None
        assert any(v == "GRAND TOTAL" for v in last_vals)

        # Currency column formatted with $
        # Find row with line totals (data row after header)
        for row in ws.iter_rows(min_row=header_row_idx + 1, max_row=header_row_idx + 30):
            if isinstance(row[6].value, (int, float)) and row[6].value > 0:
                assert "$" in (row[6].number_format or "")
                break

        # Utilities present
        all_text = " ".join(str(c.value) for row in ws.iter_rows(values_only=False) for c in row if c.value)
        assert ("Septic" in all_text or "EMT" in all_text or "Water main" in all_text.lower() or "water main" in all_text.lower())


# ---------- Share link feature ----------
class TestShare:
    def test_enable_get_rotate_disable_share(self, base_url, auth_headers_a):
        proj = _make_project(base_url, auth_headers_a)
        pid = proj["id"]

        # Enable
        r = requests.post(f"{base_url}/api/projects/{pid}/share",
                          headers=auth_headers_a, timeout=20)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["enabled"] is True
        token1 = data["token"]
        assert isinstance(token1, str) and len(token1) > 10

        # GET returns same token (idempotent)
        r = requests.get(f"{base_url}/api/projects/{pid}/share",
                         headers=auth_headers_a, timeout=20)
        assert r.status_code == 200
        assert r.json()["token"] == token1
        assert r.json()["enabled"] is True

        # Public resolve works with NO auth
        r = requests.get(f"{base_url}/api/share/{token1}", timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        for k in ("project", "owner", "blueprint", "materials", "grand_total"):
            assert k in body, f"missing key {k}"
        assert body["owner"].get("email")
        assert body["owner"].get("name") is not None
        # Sensitive fields not exposed
        assert "password_hash" not in body["owner"]
        assert "entitlements" not in body["owner"]
        # materials list always present
        assert isinstance(body["materials"], list)

        # Public CSV w/ no auth
        r = requests.get(f"{base_url}/api/share/{token1}/takeoff.csv", timeout=20)
        assert r.status_code == 200
        assert "text/csv" in r.headers.get("content-type", "")
        assert r.content[:3] == b"\xef\xbb\xbf"

        # Rotate token
        r = requests.post(f"{base_url}/api/projects/{pid}/share/rotate",
                          headers=auth_headers_a, timeout=20)
        assert r.status_code == 200
        token2 = r.json()["token"]
        assert token2 != token1

        # Old token now 404
        r_old = requests.get(f"{base_url}/api/share/{token1}", timeout=20)
        assert r_old.status_code == 404

        # New token resolves
        r_new = requests.get(f"{base_url}/api/share/{token2}", timeout=20)
        assert r_new.status_code == 200

        # Disable
        r = requests.delete(f"{base_url}/api/projects/{pid}/share",
                            headers=auth_headers_a, timeout=20)
        assert r.status_code == 200
        assert r.json()["enabled"] is False

        # Token no longer resolves
        r = requests.get(f"{base_url}/api/share/{token2}", timeout=20)
        assert r.status_code == 404

    def test_share_isolation_user_b_cannot_access(self, base_url, auth_headers_a, auth_headers_b):
        # A creates project + share
        proj = _make_project(base_url, auth_headers_a)
        pid = proj["id"]
        r = requests.post(f"{base_url}/api/projects/{pid}/share",
                          headers=auth_headers_a, timeout=20)
        assert r.status_code == 200
        # B tries to read A's share metadata -> 404
        r = requests.get(f"{base_url}/api/projects/{pid}/share",
                         headers=auth_headers_b, timeout=20)
        assert r.status_code == 404

    def test_bogus_token_returns_404(self, base_url):
        r = requests.get(f"{base_url}/api/share/nonexistent_token_xyz", timeout=20)
        assert r.status_code == 404
